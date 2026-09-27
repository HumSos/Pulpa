from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = "Arial"
MONEDA = '"$"#,##0.00'
FECHA = "DD/MM/YYYY"


def libro(hoja: str, encabezados: list[str], filas: list[list]) -> bytes:
    """Una hoja con encabezado fijo, anchos ajustados y formato por tipo de dato."""
    wb = Workbook()
    ws = wb.active
    ws.title = hoja[:31]

    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1C5240")
        celda.alignment = Alignment(vertical="center")

    for fila in filas:
        ws.append(fila)

    anchos = [len(str(e)) for e in encabezados]
    for fila in ws.iter_rows(min_row=2):
        for indice, celda in enumerate(fila):
            celda.font = Font(name=FUENTE)
            valor = celda.value
            if isinstance(valor, Decimal | float):
                celda.number_format = MONEDA
                celda.alignment = Alignment(horizontal="right")
            elif isinstance(valor, datetime | date):
                celda.number_format = FECHA
            anchos[indice] = max(anchos[indice], len(str(valor or "")))

    for indice, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = min(ancho + 3, 50)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    memoria = BytesIO()
    wb.save(memoria)
    return memoria.getvalue()


def cuentas_por_cobrar(pedidos, hoy: date) -> bytes:
    filas = [
        [
            p.folio,
            p.cliente.nombre_negocio,
            p.cliente.plaza or "",
            p.fecha_pedido.date(),
            p.fecha_vencimiento,
            max((hoy - p.fecha_vencimiento).days, 0),
            p.total,
            p.cobrado,
            p.saldo,
        ]
        for p in pedidos
    ]
    return libro(
        "Por cobrar",
        ["Pedido", "Cliente", "Plaza", "Fecha", "Vence", "Días vencido",
         "Total", "Cobrado", "Saldo"],
        filas,
    )


def cuentas_por_pagar(ordenes, hoy: date) -> bytes:
    filas = [
        [
            o.folio,
            o.proveedor.nombre,
            o.numero_remision or "",
            o.fecha.date(),
            o.fecha_vencimiento,
            max((hoy - o.fecha_vencimiento).days, 0),
            o.total,
            o.pagado,
            o.saldo,
        ]
        for o in ordenes
    ]
    return libro(
        "Por pagar",
        ["Orden", "Proveedor", "Remisión", "Fecha", "Vence", "Días vencido",
         "Total", "Pagado", "Saldo"],
        filas,
    )


def ventas_detalladas(pedidos) -> bytes:
    """Una fila por línea de pedido, que es como lo llevan en su archivo."""
    filas = []
    for pedido in pedidos:
        for linea in pedido.lineas:
            filas.append(
                [
                    pedido.folio,
                    pedido.fecha_pedido.date(),
                    pedido.cliente.nombre_negocio,
                    pedido.cliente.tipo_negocio or "",
                    pedido.cliente.plaza or "",
                    pedido.vendedor or "",
                    pedido.estado.value,
                    linea.descripcion,
                    linea.unidad_clave,
                    float(linea.cantidad),
                    linea.precio_unitario,
                    linea.importe,
                ]
            )
    return libro(
        "Ventas",
        ["Pedido", "Fecha", "Cliente", "Giro", "Plaza", "Vendedor", "Estado",
         "Producto", "UM", "Cantidad", "P. unitario", "Importe"],
        filas,
    )