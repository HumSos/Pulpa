from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.modelos import Producto, Unidad
from app.servicios.catalogo import registrar_precio
from app.servicios.clientes import DatosCliente, crear_cliente

ENERO = datetime(2026, 1, 1)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def datos(db):
    kg = Unidad(clave="kg", nombre="Kilogramo", permite_decimales=True)
    mango = Producto(codigo="PUL-001", descripcion="Pulpa de mango", unidad=kg)
    db.add(mango)
    db.flush()
    registrar_precio(db, mango.id, Decimal("100.00"), ENERO)
    cliente = crear_cliente(
        db, DatosCliente("Jugos Doña Mary", dias_credito=15, plaza="Saltillo")
    )
    db.commit()
    return {"mango": mango, "cliente": cliente}


def _pedido(client, datos, cantidad="10", confirmar=True):
    creado = client.post(
        "/pedidos/nuevo",
        data={"cliente_id": datos["cliente"].id, "vendedor": "GI"},
        follow_redirects=False,
    )
    url = creado.headers["location"]
    client.post(f"{url}/lineas", data={"producto_id": datos["mango"].id, "cantidad": cantidad})
    if confirmar:
        client.post(f"{url}/estado", data={"estado": "confirmado"})
    return url


def test_la_nota_de_entrega_muestra_lineas_y_total(client, datos):
    url = _pedido(client, datos)
    nota = client.get(f"{url}/nota")

    assert nota.status_code == 200
    assert "NOTA DE ENTREGA" in nota.text
    assert "Pulpa de mango" in nota.text
    assert "$1,000.00" in nota.text
    assert "Recibió de conformidad" in nota.text


def test_el_estado_de_cuenta_suma_saldos(client, datos):
    _pedido(client, datos, "3")
    _pedido(client, datos, "5")

    estado = client.get(f"/clientes/{datos['cliente'].id}/estado-cuenta")
    assert "ESTADO DE CUENTA" in estado.text
    assert "$800.00" in estado.text


def test_estado_de_cuenta_de_cliente_inexistente(client):
    assert client.get("/clientes/999/estado-cuenta").status_code == 404


def test_excel_por_cobrar_trae_encabezados_y_datos(client, datos):
    _pedido(client, datos, "4")
    respuesta = client.get("/cobranza/por-cobrar.xlsx")

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == XLSX
    assert "por_cobrar_" in respuesta.headers["content-disposition"]

    wb = load_workbook(BytesIO(respuesta.content))
    ws = wb.active
    encabezados = [celda.value for celda in ws[1]]
    assert encabezados[:3] == ["Pedido", "Cliente", "Plaza"]
    assert ws.cell(row=2, column=2).value == "Jugos Doña Mary"
    assert ws.cell(row=2, column=9).value == 400


def test_excel_de_ventas_una_fila_por_linea(client, datos):
    _pedido(client, datos, "2")
    _pedido(client, datos, "3", confirmar=False)

    wb = load_workbook(BytesIO(client.get("/pedidos.xlsx").content))
    ws = wb.active
    assert ws.max_row == 3  # encabezado más dos líneas
    assert "Vendedor" in [celda.value for celda in ws[1]]